"""Estilos: responsabilidad separada de gestor/servicios/excel.py."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AREA = {
    'gestion_social': 'Gestión Social',
    'atencion_ciudadano': 'Atención al Ciudadano',
    'comunicaciones': 'Comunicaciones',
}

DIAS = {
    0: 'Lunes', 1: 'Martes', 2: 'Miércoles', 3: 'Jueves',
    4: 'Viernes', 5: 'Sábado', 6: 'Domingo',
}

MESES = [
    'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
    'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre',
]

MESES_CORTOS = ['ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov', 'dic']

MARCA_AGUA = 'Creado por xYojanSaidx'

HOJA_SEMANAS = 'Semanas completas'

HOJA_MES = 'Solo el mes'

FRANJAS = [
    ['Código', 'Tipo', 'Aplicación', 'Horario', 'Descanso', 'Horas efectivas'],
    ['AM', 'Operativo AM', 'Lunes a sábado', '5:00 a. m. a 1:00 p. m.', '1 hora', 7],
    ['AM', 'Operativo AM', 'Domingo o festivo', '5:30 a. m. a 1:30 p. m.', '1 hora', 7],
    ['PM', 'Operativo PM', 'Lunes a sábado', '1:00 p. m. a 9:00 p. m.', '1 hora', 7],
    ['PM', 'Operativo PM', 'Domingo o festivo', '12:00 m. a 8:00 p. m.', '1 hora', 7],
    ['ADM-GS', 'Administrativo Guía Social', 'Lunes a domingo', '7:30 a. m. a 3:30 p. m.', '1 hora', 7],
    ['ADM-AC', 'Administrativo AC', 'Lunes a viernes', '7:30 a. m. a 4:00 p. m.', '1 hora', 7.5],
    ['ADM-AC', 'Administrativo AC', 'Sábado', '7:30 a. m. a 12:00 m.', 'Sin pausa', 4.5],
    ['CAP', 'Capacitación', 'Solicitud aprobada', 'Horario solicitado', 'Según capacitación', 'Horas reales'],
    ['D', 'Descanso', 'Día asignado', 'No aplica', 'No aplica', 0],
    ['VAC', 'Vacaciones', 'Novedad aprobada', 'No aplica', 'No aplica', 0],
    ['INC', 'Incapacidad', 'Novedad aprobada', 'No aplica', 'No aplica', 0],
    ['PER', 'Permiso', 'Novedad aprobada', 'No aplica', 'No aplica', 0],
]

CODIGOS_NO_LABORADOS = ['D', 'VAC', 'INC', 'PER', 'NV']

CODIGOS_TRABAJO = ['AM', 'PM', 'ADM-GS', 'ADM-AC', 'CAP']

EMPIEZA_UNA_FORMULA = ('=', '+', '-', '@', '\t', '\r')

LIBROS = {
    'trabajo': '_Horario',
    'semana': '_Semana',
    'completo': '_Completa',
}


def _borde():
    side = Side(style='thin', color='B7B7B7')
    return Border(left=side, right=side, top=side, bottom=side)

def _header(cell, color='5B2A86'):
    cell.fill = PatternFill('solid', fgColor=color)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = _borde()

def _cell(cell, center=True):
    cell.border = _borde()
    cell.alignment = Alignment(
        horizontal='center' if center else 'left',
        vertical='center',
        wrap_text=True,
    )

def _solo_texto(cell):
    """Lo que alguien escribió se imprime; no se ejecuta.

    Los textos libres del horario —el cargo, el nombre, la descripción de una
    asignación— se entregaban a openpyxl tal cual. Una persona apuntada como
    `=1+1` se guardaba como **fórmula**, y el Excel que se reparte enseñaba `2`
    donde tenía que ir un nombre. Con algo menos inocente que `=1+1` —y las
    hojas de cálculo tienen funciones que leen archivos y abren enlaces— el
    papel deja de ser un papel.

    Marcar la celda como cadena basta en un .xlsx: se guarda entre los textos
    del libro y Excel no vuelve a mirarla como fórmula.

    No se hace dentro de `_cell` a propósito: el libro tiene celdas que **sí**
    son fórmulas puestas por el programa —el espejo que trae el turno desde la
    hoja de semanas, y los contadores auxiliares de rachas— y pasan por ahí.
    Esto se aplica a los textos que vienen de la plantilla, que son los que
    escribe una persona.
    """
    valor = cell.value
    if isinstance(valor, str) and valor.startswith(EMPIEZA_UNA_FORMULA):
        cell.data_type = 's'
    return cell

def _fecha_corta(iso: str) -> str:
    _a, m, d = str(iso).split('-')
    return f'{int(d)} {MESES_CORTOS[int(m) - 1]}'

def _dow(d: dict) -> int:
    """Día de la semana del registro, calculándolo si no viene dado."""
    valor = d.get('dia_semana_numero')
    if valor is not None:
        return int(valor)
    from datetime import date as _date
    return _date.fromisoformat(str(d['fecha'])).weekday()

def _agrupar_semanas(dias: list[dict]) -> list[list[dict]]:
    """Parte los días del período en sus semanas de lunes a domingo."""
    semanas: list[list[dict]] = []
    from datetime import date as _date
    from datetime import timedelta as _td
    for d in dias:
        lunes = d.get('lunes_semana')
        if not lunes:
            f = _date.fromisoformat(str(d['fecha']))
            lunes = (f - _td(days=f.weekday())).isoformat()
            d['lunes_semana'] = lunes
        if semanas and semanas[-1][0].get('lunes_semana') == lunes:
            semanas[-1].append(d)
        else:
            semanas.append([d])
    return semanas

def _rangos(columnas: list[int], fila: int) -> str:
    """Convierte columnas sueltas en la lista de rangos contiguos de una fila."""
    if not columnas:
        return ''
    partes = []
    inicio = anterior = columnas[0]
    for col in columnas[1:]:
        if col == anterior + 1:
            anterior = col
            continue
        partes.append((inicio, anterior))
        inicio = anterior = col
    partes.append((inicio, anterior))
    return ','.join(
        f'{get_column_letter(a)}{fila}' if a == b
        else f'{get_column_letter(a)}{fila}:{get_column_letter(b)}{fila}'
        for a, b in partes
    )

def _contar(rangos: str, codigos: list[str]) -> str:
    if not rangos:
        return '0'
    piezas = [f'COUNTIF({r},"{c}")' for r in rangos.split(',') for c in codigos]
    return '+'.join(piezas) if piezas else '0'

def _formula_horas(fila: int, todos: list[int], ac_semana: list[int], ac_sabado: list[int]) -> str:
    """Horas efectivas del rango, con las franjas reales de cada código."""
    r_todos = _rangos(todos, fila)
    partes = [f'7*({_contar(r_todos, ["AM", "PM", "ADM-GS", "CAP"])})']
    r_semana = _rangos(ac_semana, fila)
    if r_semana:
        partes.append(f'7.5*({_contar(r_semana, ["ADM-AC"])})')
    r_sabado = _rangos(ac_sabado, fila)
    if r_sabado:
        partes.append(f'4.5*({_contar(r_sabado, ["ADM-AC"])})')
    return '=' + '+'.join(partes)
