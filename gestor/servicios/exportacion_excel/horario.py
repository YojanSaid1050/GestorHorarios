"""Horario: responsabilidad separada de gestor/servicios/excel.py."""
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gestor.registro import obtener as obtener_registro
from gestor.servicios.exportacion_excel.estilos import (
    AREA,
    CODIGOS_NO_LABORADOS,
    CODIGOS_TRABAJO,
    DIAS,
    FRANJAS,
    MARCA_AGUA,
    MESES,
    MESES_CORTOS,
    _agrupar_semanas,
    _cell,
    _contar,
    _dow,
    _fecha_corta,
    _formula_horas,
    _header,
    _rangos,
    _solo_texto,
)
from gestor.servicios.reglas_operacion import POR_DEFECTO as MAX_DIAS_POR_DEFECTO


def _hoja_horario(wb, titulo: str, resultado: dict, dias: list[dict],
                  colores: dict, opciones: dict) -> dict:
    """Construye una hoja de horario y devuelve lo que necesitan las demás.

    Una sola función para las tres formas que puede tomar la hoja:

    * la de trabajo, con todas las columnas de control en fórmula;
    * su espejo del mes, cuyos turnos son una referencia a la hoja de
      semanas: se cambia un turno allí y aquí cambia solo;
    * la informativa, que se queda con el horario y las dos cifras de horas.
    """
    con_comentarios = bool(opciones.get('con_comentarios'))
    controles = bool(opciones.get('controles', True))
    enlace = opciones.get('enlace')
    enlace_cols = opciones.get('enlace_cols') or {}
    color_header = colores['header']
    color_sabado = colores['saturday_header']
    color_domingo = colores['sunday_header']
    color_festivo = colores['holiday_header']
    color_viernes_adm = colores['admin_friday_header']
    color_celda_viernes_adm = colores['admin_friday_cell']
    color_otro_mes = colores.get('otro_mes_header', '41648C')
    color_otro_mes_celda = colores.get('otro_mes_cell', 'EDF3F9')

    ws = wb.create_sheet(titulo)
    horario = resultado['horario']
    semanas = _agrupar_semanas(dias)
    mes = int(resultado.get('mes') or 0)
    anio = int(resultado.get('anio') or 0)
    nombre_mes = MESES[mes - 1] if 1 <= mes <= 12 else ''

    nombres_por_id = {int(e['empleado_id']): e['nombre'] for e in horario}

    # ------------------------------------------------------------------
    # Geometría de la hoja
    # ------------------------------------------------------------------
    base_headers = ['Cargo', 'Nombre del colaborador', 'Área', 'Turno base', 'Descanso obligatorio', 'Pareja']
    start_col = len(base_headers) + 1
    end_day_col = start_col + len(dias) - 1
    stats_start = end_day_col + 1

    fechas_hoja = [d['fecha'] for d in dias]
    col_por_fecha = {d['fecha']: start_col + i for i, d in enumerate(dias)}
    cols_mes = [col_por_fecha[d['fecha']] for d in dias if d.get('mes_propio', True)]
    # El período son todas las columnas, incluidos los días que solo completan
    # la primera y la última semana. Esa es la cifra que queda pareja.
    cols_periodo = [col_por_fecha[d['fecha']] for d in dias]
    cols_periodo_ac_semana = [col_por_fecha[d['fecha']] for d in dias if _dow(d) <= 4]
    cols_periodo_ac_sabado = [col_por_fecha[d['fecha']] for d in dias if _dow(d) == 5]
    cols_mes_ac_semana = [
        col_por_fecha[d['fecha']] for d in dias
        if d.get('mes_propio', True) and _dow(d) <= 4
    ]
    cols_mes_ac_sabado = [
        col_por_fecha[d['fecha']] for d in dias
        if d.get('mes_propio', True) and _dow(d) == 5
    ]
    cols_domingos_mes = [
        col_por_fecha[d['fecha']] for d in dias
        if d.get('mes_propio', True) and d['es_domingo']
    ]
    hay_dias_de_otro_mes = any(not d.get('mes_propio', True) for d in dias)
    # Una hoja va «recortada» cuando no contiene el período entero: se pidió
    # solo el mes natural o una sola semana.
    hoja_recortada = bool(opciones.get('recorte'))
    recorte_texto = str(opciones.get('recorte') or '')

    # Dos cifras de horas, porque dicen cosas distintas y las dos importan:
    # las del período son las que quedan parejas para todo el personal —es lo
    # que el horario equilibra— y las del mes natural varían según dónde caiga
    # el descanso de cada quien, que es normal y no es un desequilibrio.
    reglas = resultado.get('reglas', {})
    # El tope de jornadas seguidas es configurable en la aplicación. Si el
    # horario lo trae anotado se usa ese —es el que regía cuando se generó—; si
    # no, se pregunta por el que rige ahora. Escribirlo a mano en la fórmula
    # haría que el Excel marcara «REVISAR» en rachas perfectamente legales.
    tope_racha = reglas.get('maximo_dias_consecutivos')
    if not tope_racha:
        # Esto apuntaba a `backend.operation_rules`, que es el paquete de la
        # versión anterior y aquí no existe: el import fallaba siempre, el
        # `except` se lo tragaba y el Excel salía impreso diciendo «máximo 7
        # jornadas seguidas» mientras la regla configurada era 10. Un fallo que
        # no daba error, no salía en ningún registro y viajaba en el papel que
        # se reparte.
        try:
            from gestor.servicios.reglas_operacion import maximo_dias
            tope_racha = maximo_dias(dias[0]['fecha']) if dias else None
        except Exception:                                          # noqa: BLE001
            obtener_registro().exception(
                'no se pudo leer el máximo de jornadas seguidas para el Excel')
            tope_racha = None
    tope_racha = int(tope_racha or MAX_DIAS_POR_DEFECTO)

    headers_stats = ['Horas del período (dato)' if hoja_recortada else 'Horas del período',
                     f'Horas {nombre_mes}' if nombre_mes else 'Horas del mes',
                     'Días trabajados',
                     'Domingos trabajados',
                     'Domingos descansados',
                     'Objetivo domingos',
                     'Control domingos']
    if not controles:
        # Los libros informativos se quedan con el horario y las horas: los
        # controles son de quien programa, no de quien lo lee.
        headers_stats = headers_stats[:2]
    headers_semana = ([f'Horas sem. {_fecha_corta(s[0]["fecha"])}–{_fecha_corta(s[-1]["fecha"])}' for s in semanas]
                      if controles else [])
    headers_final = (['Máx. días seguidos', f'Control fatiga (máx. {tope_racha})', 'Observaciones']
                     if controles else ['Observaciones'])
    todas_stats = headers_stats + headers_semana + headers_final
    end_col = stats_start + len(todas_stats) - 1

    domingos = reglas.get('domingos_mes', len(cols_domingos_mes))
    objetivo = reglas.get('domingos_trabajo_objetivo', domingos // 2)
    objetivo_max = reglas.get('domingos_trabajo_maximo_equilibrado', objetivo + (domingos % 2))

    # ------------------------------------------------------------------
    # Cabecera
    # ------------------------------------------------------------------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    if resultado.get('_export_week_start'):
        alcance = 'PROGRAMACIÓN SEMANAL DE TURNOS'
    elif hoja_recortada:
        alcance = 'PROGRAMACIÓN MENSUAL DE TURNOS'
    else:
        alcance = 'PROGRAMACIÓN DE TURNOS POR SEMANAS COMPLETAS'
    titulo_texto = f'{alcance} · {nombre_mes.upper()} {anio}' if nombre_mes else alcance
    if resultado.get('_es_borrador_export'):
        titulo_texto = f'BORRADOR — NO PUBLICADO · {titulo_texto}'
    title = ws.cell(1, 1, titulo_texto)
    title.font = Font(size=16, bold=True, color='FFFFFF')
    title.fill = PatternFill('solid', fgColor=color_header)
    title.alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    if hay_dias_de_otro_mes:
        subtitulo = (
            f'Vista semanal completa del {_fecha_corta(dias[0]["fecha"])} al {_fecha_corta(dias[-1]["fecha"])}. '
            f'Los días fuera de {nombre_mes} completan la primera y la última semana: sirven para comprobar '
            f'descansos, rachas y continuidad, y no suman a «Horas {nombre_mes}».'
        )
    else:
        subtitulo = (
            f'Del {_fecha_corta(dias[0]["fecha"])} al {_fecha_corta(dias[-1]["fecha"])}. '
            'AM y PM son los turnos operativos; D es descanso; ADM-GS y ADM-AC, jornadas administrativas.'
        )
    ws.cell(2, 1, subtitulo)
    ws.cell(2, 1).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.cell(2, 1).font = Font(italic=True)

    # Fila 3: bandas de semana
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(base_headers))
    _header(ws.cell(3, 1, f'Período oficial: {nombre_mes} de {anio}' if nombre_mes else 'Período oficial'), color_header)
    for indice, semana in enumerate(semanas, 1):
        primera = col_por_fecha[semana[0]['fecha']]
        ultima = col_por_fecha[semana[-1]['fecha']]
        etiqueta = f'Semana {indice} · {_fecha_corta(semana[0]["fecha"])}–{_fecha_corta(semana[-1]["fecha"])}'
        if primera != ultima:
            ws.merge_cells(start_row=3, start_column=primera, end_row=3, end_column=ultima)
        _header(ws.cell(3, primera, etiqueta), color_header)
    ws.merge_cells(start_row=3, start_column=stats_start, end_row=3, end_column=end_col)
    _header(ws.cell(3, stats_start, 'Controles y horas · fórmulas que se recalculan solas'), color_header)

    # Fila 4: encabezados. Fila 5: día de la semana.
    for col, text in enumerate(base_headers, 1):
        _header(ws.cell(4, col, text), color_header)
        _header(ws.cell(5, col, ''), color_header)

    for col, day in enumerate(dias, start_col):
        otro_mes = not day.get('mes_propio', True)
        color = (
            color_otro_mes if otro_mes else
            color_festivo if day['es_festivo'] else
            color_viernes_adm if day.get('es_ultimo_viernes_administrativo') else
            color_domingo if day['es_domingo'] else
            color_sabado if day.get('es_sabado') else
            color_header
        )
        etiqueta = f"{day['dia']} {MESES_CORTOS[int(day.get('mes') or mes) - 1]}" if otro_mes else day['dia']
        day_cell = ws.cell(4, col, etiqueta)
        _header(day_cell, color)
        _header(ws.cell(5, col, day['dia_semana'][:3]), color)
        notas = []
        if otro_mes:
            notas.append('Día de otro mes: completa la semana y no suma a las horas del mes.')
        if day.get('nombre_festivo'):
            notas.append(f"Festivo: {day['nombre_festivo']}")
        elif day.get('es_ultimo_viernes_administrativo'):
            notas.append('Último viernes administrativo del mes')
        if notas and con_comentarios:
            # «Sin comentarios» es sin ninguno, también en la fila de fechas: el
            # día de otro mes ya se distingue por su color y su marco.
            day_cell.comment = Comment('\n'.join(notas), 'Sistema')

    for col, text in enumerate(todas_stats, stats_start):
        _header(ws.cell(4, col, text), color_header)
        _header(ws.cell(5, col, ''), color_header)

    # ------------------------------------------------------------------
    # Filas de personal
    # ------------------------------------------------------------------
    stats_by = {s['empleado_id']: s for s in resultado.get('estadisticas', [])}
    primera_fila = 6
    ultima_fila = primera_fila + len(horario) - 1

    notas = [
        f'Regla de domingos: {nombre_mes} de {anio} tiene {domingos} domingos; el personal operativo trabaja '
        f'{objetivo} y descansa {domingos - objetivo}. El horario administrativo conserva domingos y festivos no laborados.',
        'Las horas son informativas y reflejan el tiempo programado. Una semana ordinaria suele sumar 42 h; '
        'un festivo, un compensatorio o una ausencia aprobada pueden reducir el total sin que eso sea un conflicto.',
        'Si editas una celda de turno dentro de este Excel, las horas, los domingos y el máximo de días seguidos '
        'se recalculan solos. El cambio no vuelve a la aplicación: para que quede guardado hazlo desde la app.',
        f'Control de fatiga: el tope configurado es de {tope_racha} jornadas seguidas. La columna marca REVISAR '
        f'a partir de la {tope_racha + 1}.',
    ]
    if hoja_recortada:
        notas.append(
            f'«Horas del período (dato)» es la cifra del período completo, de semanas enteras: la que el '
            f'horario equilibra y en la que todo el personal queda parejo. Aquí va como dato y no como '
            f'fórmula porque esta hoja solo trae {recorte_texto}, y los días que faltan no están para poder '
            f'sumarlos. Si la necesitas calculada, exporta el libro del período.'
        )
    if hay_dias_de_otro_mes:
        fuera = [d for d in dias if not d.get('mes_propio', True)]
        etiquetas = ', '.join(_fecha_corta(d['fecha']) for d in fuera)
        notas.append(
            f'Las columnas {etiquetas} completan la primera y la última semana. Sirven para verificar horas '
            f'semanales, descansos, fatiga y continuidad; no se incluyen en «Horas {nombre_mes}».'
        )

    fila_nota = ultima_fila + 2
    fila_franjas = fila_nota + len(notas) + 1
    fila_marca = fila_franjas + len(FRANJAS) + 2
    # Filas auxiliares: cuentan días seguidos de trabajo. Van al final y se
    # ocultan porque son el andamiaje de la fórmula «Máx. días seguidos», no
    # información para leer.
    fila_aux_inicio = fila_marca + 3

    for indice, employee in enumerate(horario):
        row = primera_fila + indice
        fila_aux = fila_aux_inicio + indice
        stat = stats_by.get(employee['empleado_id'], {})

        if employee['tipo_turno'] == 'administrativo':
            descanso = 'Domingos y festivos'
        elif employee.get('descanso_fijo') is not None:
            rango = str(objetivo) if objetivo == objetivo_max else f'{objetivo}-{objetivo_max}'
            descanso = f"{DIAS[employee['descanso_fijo']]} fijo + domingos {rango}/{domingos}"
        else:
            rango = str(objetivo) if objetivo == objetivo_max else f'{objetivo}-{objetivo_max}'
            descanso = f'Variable · domingos {rango}/{domingos}'

        pareja = nombres_por_id.get(int(employee.get('pareja_id') or 0), '—')

        values = [
            employee.get('cargo', 'GUÍA SOCIAL'),
            employee['nombre'],
            AREA.get(employee['area'], employee['area']),
            employee['turno_base'],
            descanso,
            pareja,
        ]
        for col, value in enumerate(values, 1):
            _cell(_solo_texto(ws.cell(row, col, value)), center=col != 2)

        # Los días de esta hoja, no los de todo el período: la hoja del mes
        # tiene menos columnas y antes se le escribían de más, corridas.
        dias_empleado = {d['fecha']: d for d in employee['dias']}
        for col, fecha_dia in enumerate(fechas_hoja, start_col):
            day = dias_empleado.get(fecha_dia)
            if day is None:
                continue
            otro_mes = not day.get('mes_propio', True)
            if enlace and day['fecha'] in enlace_cols:
                # Espejo: el turno no se copia, se toma de la hoja de semanas.
                # Así un cambio hecho allí aparece aquí sin tocar nada.
                origen = f"'{enlace}'!{get_column_letter(enlace_cols[day['fecha']])}{row}"
                cell = ws.cell(row, col, f'={origen}')
            else:
                cell = ws.cell(row, col, '' if day.get('turno') == 'NV' else day['turno'])
            _cell(cell)
            if otro_mes and day['turno'] in {'', 'NV'}:
                color_celda = color_otro_mes_celda
            elif day.get('es_ultimo_viernes_administrativo') and day['turno'] in {'ADM-GS', 'ADM-AC'}:
                color_celda = color_celda_viernes_adm
            else:
                color_celda = colores.get(day['turno'], 'FFFFFF')
            cell.fill = PatternFill('solid', fgColor=color_celda)
            cell.font = Font(bold=True)
            if otro_mes:
                # Marco del color de «días de otro mes» sin perder el color del turno.
                lado = Side(style='medium', color=color_otro_mes)
                cell.border = Border(left=lado, right=lado, top=lado, bottom=lado)

            detalles = [
                'Día publicado con el mes anterior · no se modifica aquí' if day.get('heredado') else None,
                'Completa la semana; pertenece a otro mes' if otro_mes and not day.get('heredado') else None,
                day.get('observacion'),
                f"Festivo: {day['nombre_festivo']}" if day.get('nombre_festivo') else None,
                f"Capacitación: {day.get('capacitacion_horas', 0):.1f} h" if day.get('capacitacion_horas') else None,
            ]
            detalles = [x for x in detalles if x]
            if detalles and con_comentarios:
                cell.comment = Comment('\n'.join(detalles), 'Sistema')

        # --- Controles con fórmula ------------------------------------
        col = stats_start
        rango_mes = _rangos(cols_mes, row)
        rango_dom = _rangos(cols_domingos_mes, row)

        if enlace and opciones.get('enlace_horas_periodo'):
            # También el total del período viene de la hoja de semanas: si allí
            # cambia un turno, aquí cambia la cifra. Es la única forma de que la
            # columna diga la verdad en una hoja que no tiene esos días.
            letra = get_column_letter(int(opciones['enlace_horas_periodo']))
            _cell(ws.cell(row, col, f"='{enlace}'!{letra}{row}"))
        elif hoja_recortada:
            # La hoja no trae todos los días del período —se pidió solo el mes o
            # una semana—, así que aquí no hay con qué calcularlo. Antes se
            # ponía la fórmula igual y salía la cifra del recorte: dos columnas
            # distintas diciendo lo mismo, y la de la izquierda mintiendo. Ahora
            # se escribe el dato real, y se dice de dónde sale.
            horas_periodo = stat.get('horas_periodo')
            celda_hp = ws.cell(row, col, horas_periodo if horas_periodo is not None else '—')
            _cell(celda_hp)
            if con_comentarios:
                celda_hp.comment = Comment(
                    'Horas del período completo, de semanas enteras.\n'
                    'Esta hoja solo trae ' + recorte_texto + ', por eso no es una fórmula: '
                    'los días que faltan no están en la hoja para poder sumarlos.\n'
                    'Para verla calculada, exporta el libro del período.',
                    'Sistema')
        else:
            _cell(ws.cell(row, col, _formula_horas(
                row, cols_periodo, cols_periodo_ac_semana, cols_periodo_ac_sabado)))
        col += 1
        _cell(ws.cell(row, col, _formula_horas(row, cols_mes, cols_mes_ac_semana, cols_mes_ac_sabado)))
        col += 1
        if controles:
            _cell(ws.cell(row, col, f'={_contar(rango_mes, CODIGOS_TRABAJO)}'))
            col += 1
            _cell(ws.cell(row, col, f'={_contar(rango_dom, CODIGOS_TRABAJO)}'))
            col += 1
            _cell(ws.cell(row, col, f'={_contar(rango_dom, CODIGOS_NO_LABORADOS)}'))
            col += 1
            exento = (
                employee['tipo_turno'] == 'administrativo'
                or employee.get('descanso_fijo') is not None
                or employee.get('exento_especiales')
            )
            objetivo_txt = 'Exento' if exento else (
                f'{objetivo} trabajo / {domingos - objetivo} D' if objetivo == objetivo_max
                else f'{objetivo}-{objetivo_max} de {domingos}'
            )
            _cell(ws.cell(row, col, objetivo_txt))
            col += 1
            letra_dom = get_column_letter(stats_start + 3)
            if exento:
                _cell(ws.cell(row, col, 'EXENTO'))
            else:
                _cell(ws.cell(
                    row, col,
                    f'=IF(AND({letra_dom}{row}>={objetivo},{letra_dom}{row}<={objetivo_max}),"CUMPLE","REVISAR")',
                ))
            col += 1

        if controles:
            for semana in semanas:
                cols_s = [col_por_fecha[d['fecha']] for d in semana]
                cols_s_ac = [col_por_fecha[d['fecha']] for d in semana if _dow(d) <= 4]
                cols_s_sab = [col_por_fecha[d['fecha']] for d in semana if _dow(d) == 5]
                _cell(ws.cell(row, col, _formula_horas(row, cols_s, cols_s_ac, cols_s_sab)))
                col += 1

        if controles:
            primera_col_letra = get_column_letter(start_col)
            ultima_col_letra = get_column_letter(end_day_col)
            _cell(ws.cell(row, col, f'=MAX({primera_col_letra}{fila_aux}:{ultima_col_letra}{fila_aux})'))
            letra_max = get_column_letter(col)
            col += 1
            _cell(ws.cell(
                row, col,
                f'=IF({letra_max}{row}<={tope_racha},"CUMPLE","REVISAR")',
            ))
            col += 1
        _cell(ws.cell(row, col, ''), center=False)

        # Fila auxiliar: racha de días trabajados acumulada, día a día.
        for posicion, day_col in enumerate(range(start_col, end_day_col + 1) if controles else []):
            letra = get_column_letter(day_col)
            condiciones = ','.join(f'{letra}{row}="{c}"' for c in [''] + CODIGOS_NO_LABORADOS)
            if posicion == 0:
                ws.cell(fila_aux, day_col, f'=IF(OR({condiciones}),0,1)')
            else:
                anterior = get_column_letter(day_col - 1)
                ws.cell(fila_aux, day_col, f'=IF(OR({condiciones}),0,{anterior}{fila_aux}+1)')

    for fila in range(fila_aux_inicio, fila_aux_inicio + len(horario)):
        ws.row_dimensions[fila].hidden = True

    # ------------------------------------------------------------------
    # Notas y franjas horarias
    # ------------------------------------------------------------------
    for offset, texto in enumerate(notas):
        fila = fila_nota + offset
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=end_col)
        celda = ws.cell(fila, 1, texto)
        celda.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        celda.font = Font(size=9, italic=True, color='555555')

    rr = fila_franjas
    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
    ws.cell(rr, 1, 'HORARIOS APLICABLES')
    _header(ws.cell(rr, 1), color_header)
    for row_index, row_data in enumerate(FRANJAS, rr + 1):
        for col_index, value in enumerate(row_data, 1):
            cell = ws.cell(row_index, col_index, value)
            _header(cell, color_header) if row_index == rr + 1 else _cell(cell, center=col_index != 2)

    ws.merge_cells(start_row=fila_marca, start_column=1, end_row=fila_marca, end_column=end_col)
    marca = ws.cell(fila_marca, 1, MARCA_AGUA)
    marca.alignment = Alignment(horizontal='center', vertical='center')
    marca.font = Font(size=9, bold=True, color='999999')

    ws.freeze_panes = f'{get_column_letter(start_col)}6'
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = MARCA_AGUA
    ws.oddFooter.center.size = 8

    widths = {1: 16, 2: 31, 3: 23, 4: 13, 5: 28, 6: 28}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for col in range(start_col, end_day_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 8.5 if hay_dias_de_otro_mes else 6.5
    for col in range(stats_start, end_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 17

    # Los colores siguen al valor de la celda: si alguien cambia un turno dentro
    # del Excel —o si la hoja del mes lo recibe de la de semanas— el color
    # cambia con él. Antes el relleno se quedaba con el del turno anterior.
    _colorear_por_valor(ws, colores, start_col, end_day_col, primera_fila, ultima_fila,
                        {col_por_fecha[d['fecha']] for d in dias
                         if not d.get('mes_propio', True)
                         or (d.get('es_ultimo_viernes_administrativo')
                             and d.get('turno') in {'ADM-GS', 'ADM-AC'})})

    return {
        'ws': ws,
        'col_por_fecha': col_por_fecha,
        'col_horas_periodo': stats_start,
        'nombre_mes': nombre_mes,
        'mes': mes,
        'anio': anio,
        'domingos': domingos,
        'objetivo': objetivo,
        'tope_racha': tope_racha,
    }

def _colorear_por_valor(ws, colores: dict, primera_col: int, ultima_col: int,
                        primera_fila: int, ultima_fila: int, columnas_propias: set) -> None:
    """Hace que el color de cada celda dependa de lo que dice, no de lo que decía.

    El relleno se pinta al exportar, así que al cambiar un turno dentro del
    Excel la celda conservaba el color del turno anterior: en la hoja espejo del
    mes, que recibe sus turnos de la de semanas, eso significaba que ninguna
    celda tendría nunca el color correcto. Con formato condicional el color lo
    decide el propio valor, y son los colores configurados en la aplicación.

    Las columnas con color propio —los días de otro mes y el último viernes
    administrativo— se dejan fuera: ahí el color dice otra cosa.
    """
    if ultima_fila < primera_fila:
        return
    codigos = ['AM', 'PM', 'D', 'ADM-GS', 'ADM-AC', 'VAC', 'INC', 'PER', 'CAP']
    tramos = []
    inicio = None
    for col in range(primera_col, ultima_col + 2):
        propia = col in columnas_propias or col > ultima_col
        if propia:
            if inicio is not None:
                tramos.append((inicio, col - 1))
                inicio = None
        elif inicio is None:
            inicio = col
    for desde, hasta in tramos:
        rango = (f'{get_column_letter(desde)}{primera_fila}:'
                 f'{get_column_letter(hasta)}{ultima_fila}')
        for codigo in codigos:
            color = colores.get(codigo)
            if not color:
                continue
            ws.conditional_formatting.add(rango, CellIsRule(
                operator='equal', formula=[f'"{codigo}"'],
                fill=PatternFill('solid', start_color=color, end_color=color),
                font=Font(bold=True)))
