"""Consulta: responsabilidad separada de gestor/servicios/excel.py."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gestor.servicios.exportacion_excel.estilos import (
    AREA,
    FRANJAS,
    MARCA_AGUA,
    _cell,
    _header,
)


def _hojas_de_apoyo(wb, resultado: dict, colores: dict, ctx: dict) -> None:
    """Parámetros, Resumen mensual e Instructivo: material de consulta."""
    horario = resultado['horario']
    stats_by = {s['empleado_id']: s for s in resultado.get('estadisticas', [])}
    color_header = colores['header']
    nombre_mes = ctx['nombre_mes']
    anio = ctx['anio']

    p = wb.create_sheet('Parámetros')
    p.merge_cells('A1:F1')
    p['A1'] = 'PARÁMETROS'
    _header(p['A1'], color_header)
    for row_index, row_data in enumerate(FRANJAS, 3):
        for col_index, value in enumerate(row_data, 1):
            cell = p.cell(row_index, col_index, value)
            _header(cell, color_header) if row_index == 3 else _cell(cell, center=col_index != 2)
    for col in range(1, 7):
        p.column_dimensions[get_column_letter(col)].width = [14, 26, 25, 28, 18, 16][col - 1]

    color_row = 3 + len(FRANJAS) + 2
    p.merge_cells(start_row=color_row, start_column=1, end_row=color_row, end_column=3)
    p.cell(color_row, 1, 'COLORES DEL EXCEL')
    _header(p.cell(color_row, 1), color_header)
    color_labels = [
        ('Encabezado', 'header'), ('Sábado', 'saturday_header'), ('Domingo', 'sunday_header'),
        ('Festivo', 'holiday_header'), ('Encabezado último viernes ADM', 'admin_friday_header'),
        ('Celdas último viernes ADM', 'admin_friday_cell'),
        ('Encabezado de días de otro mes', 'otro_mes_header'),
        ('Celdas de días de otro mes', 'otro_mes_cell'),
        ('AM', 'AM'), ('PM', 'PM'), ('Descanso D', 'D'), ('ADM-GS', 'ADM-GS'),
        ('ADM-AC', 'ADM-AC'), ('Vacaciones', 'VAC'), ('Incapacidad', 'INC'),
        ('Permiso', 'PER'), ('Capacitación', 'CAP'),
    ]
    for offset, (label, key) in enumerate(color_labels, 1):
        row = color_row + offset
        p.cell(row, 1, label)
        _cell(p.cell(row, 1), center=False)
        p.cell(row, 2, f"#{colores.get(key, 'FFFFFF')}")
        _cell(p.cell(row, 2))
        p.cell(row, 3, 'Vista')
        _cell(p.cell(row, 3))
        p.cell(row, 3).fill = PatternFill('solid', fgColor=colores.get(key, 'FFFFFF'))
    p.oddFooter.center.text = MARCA_AGUA

    rs = wb.create_sheet('Resumen mensual')
    rs.merge_cells('A1:K1')
    rs['A1'] = f'RESUMEN DE {nombre_mes.upper()} {anio}' if nombre_mes else 'RESUMEN MENSUAL'
    _header(rs['A1'], color_header)
    rs.merge_cells('A2:K2')
    rs['A2'] = (
        'HORAS DEL PERÍODO: semanas completas, del primer lunes al último domingo. Es la cifra que el horario '
        'equilibra y ahí todo el personal queda parejo. · HORAS DEL MES: solo los días del mes natural; unos '
        'tendrán más y otros menos según dónde caiga su descanso, y eso es normal, porque lo que se compensa '
        'cae en el mes vecino. Por eso la programación se hace por semanas completas y no por meses. '
        'Una semana ordinaria suele sumar 42 h; ADM-GS aporta 7 h de lunes a domingo y ADM-AC 7,5 h de lunes a '
        'viernes y 4,5 h el sábado.'
    )
    rs['A2'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    rs['A2'].font = Font(italic=True, color='555555')

    summary_headers = [
        'Colaborador', 'Área', 'Horas del mes', 'Horas del período', 'Días trabajados',
        'Domingos', 'Festivos', 'Especiales trabajados', 'Objetivo especiales',
        'Control especiales', 'Control descanso',
    ]
    for col, value in enumerate(summary_headers, 1):
        _header(rs.cell(3, col, value), color_header)
    for row, employee in enumerate(horario, 4):
        stat = stats_by.get(employee['empleado_id'], {})
        values = [
            employee['nombre'],
            AREA.get(employee['area'], employee['area']),
            stat.get('horas_mes', 0),
            stat.get('horas_periodo', stat.get('horas_mes', 0)),
            stat.get('dias_trabajados', 0),
            stat.get('domingos_trabajados', 0),
            stat.get('festivos_trabajados', 0),
            stat.get('especiales_trabajados', 0),
            stat.get('objetivo_especiales', ''),
            stat.get('control_especiales', ''),
            stat.get('control_descanso', ''),
        ]
        for col, value in enumerate(values, 1):
            _cell(rs.cell(row, col, value), center=col != 1)
    for col in range(1, 12):
        rs.column_dimensions[get_column_letter(col)].width = 21
    rs.oddFooter.center.text = MARCA_AGUA

    ins = wb.create_sheet('Instructivo')
    ins.merge_cells('A1:B1')
    ins['A1'] = 'INSTRUCTIVO'
    _header(ins['A1'], color_header)
    instrucciones = [
        [1, 'El horario se lee por semanas completas de lunes a domingo. La fila de bandas indica dónde empieza y termina cada semana.'],
        [2, 'AM y PM son los turnos operativos; D, el descanso; ADM-GS y ADM-AC, las jornadas administrativas; VAC, INC, PER y CAP, las novedades aprobadas.'],
        [3, 'Los días con marco y color propio pertenecen al mes anterior o al siguiente: completan la semana y no suman a las horas del mes.'],
        [4, 'Los días heredados del mes anterior ya están publicados. Si necesitas cambiarlos, hazlo en ese mes y vuelve a generar este.'],
        [5, 'Este libro es para consultar: trae el horario, los comentarios de cada día y las dos cifras de horas, que sí se recalculan. Los controles de domingos y de fatiga están en el libro «Para trabajar».'],
        [6, 'El libro «Para trabajar» lleva dos hojas: «Semanas completas», donde se hacen los cambios, y «Solo el mes», que los recoge sola.'],
        [7, 'La hoja Parámetros recoge las franjas horarias reales y los colores configurados en la aplicación.'],
        [8, 'Un cambio hecho dentro del Excel no vuelve a la aplicación. Para que quede guardado, hazlo desde la app y exporta de nuevo.'],
    ]
    _header(ins.cell(3, 1, 'Paso'), color_header)
    _header(ins.cell(3, 2, 'Instrucción'), color_header)
    for row_index, row_data in enumerate(instrucciones, 4):
        for col_index, value in enumerate(row_data, 1):
            cell = ins.cell(row_index, col_index, value)
            _cell(cell, center=col_index == 1)
    ins.column_dimensions['A'].width = 10
    ins.column_dimensions['B'].width = 120
    ins.oddFooter.center.text = MARCA_AGUA
