# -*- coding: utf-8 -*-
"""Protecciones que dependían de por dónde se entrara, y papeles que mienten.

Una protección escrita en una de las dos rutas que hacen lo mismo no es una
protección: es una casualidad. Y un Excel que se reparte no puede llevar
impreso un mes distinto del que tiene dentro, ni ejecutar lo que alguien
escribió en un campo de texto.
"""
from __future__ import annotations

import pytest
from fastapi.testclient import TestClient

from gestor.datos import horarios


@pytest.fixture
def cliente(carpeta_de_datos):
    from gestor.web.aplicacion import crear_aplicacion
    with TestClient(crear_aplicacion()) as cliente:
        respuesta = cliente.post('/api/auth/login', json={
            'usuario': 'admin', 'password': 'xYojanSaidx1050'})
        cliente.headers['X-Session-Token'] = respuesta.json()['token']
        yield cliente


def _guardar(anio: int, mes: int, publicado: bool = False) -> int:
    import json

    from gestor.datos.base import transaccion
    with transaccion() as conexion:
        cursor = conexion.execute(
            'INSERT INTO horarios(anio, mes, datos_json, valido, grupo_id, oficial, '
            'publicado) VALUES(?,?,?,1,?,1,?)',
            (anio, mes, json.dumps({'horario': [], 'anio': anio, 'mes': mes}),
             f'{anio}-{mes:02d}-prueba', 1 if publicado else 0))
        return int(cursor.lastrowid)


# ------------------------- H19 · la protección está donde pasan los dos caminos

def test_no_se_deja_sin_oficial_un_mes_publicado_por_ninguna_de_las_dos_rutas(cliente):
    """Había dos caminos para lo mismo y solo uno miraba si estaba publicado.

    El de por año y mes llamaba directamente a `quitar_oficial`, así que por ahí
    se dejaba sin horario un mes que la oficina tenía repartido.
    """
    identificador = _guardar(2026, 10, publicado=True)

    por_id = cliente.delete(f'/api/horarios/{identificador}/oficial')
    assert por_id.status_code >= 400

    por_mes = cliente.delete('/api/horarios/oficial/2026/10')
    assert por_mes.status_code >= 400, (
        'por año y mes se dejó sin oficial un mes publicado')
    assert horarios.oficial(2026, 10) is not None


def test_un_mes_base_no_se_puede_dejar_sin_horario(cliente):
    """Agosto y septiembre llegan transcritos del Excel y no se regeneran.

    Vaciarlos no tendría vuelta atrás: no hay forma de volver a calcularlos.
    """
    from gestor.dominio import calendario

    anio, mes = calendario.MESES_BASE[0]
    _guardar(anio, mes)

    respuesta = cliente.delete(f'/api/horarios/oficial/{anio}/{mes}')
    assert respuesta.status_code >= 400
    assert 'mes base' in respuesta.text


def test_un_mes_base_tampoco_se_puede_volver_a_generar(cliente):
    from gestor.dominio import calendario

    anio, mes = calendario.MESES_BASE[0]
    respuesta = cliente.post('/api/horarios/generar', json={'mes': mes, 'anio': anio})
    assert respuesta.status_code >= 400
    assert 'transcrito' in respuesta.text


# --------------------------- H27 · el papel lleva impreso el mes que tiene dentro

def test_no_se_exporta_un_horario_con_el_nombre_de_otro_mes(cliente):
    """Se sobrescribía el período con el de la dirección, sin comparar.

    Pedir el horario de octubre desde la dirección de noviembre producía un
    Excel con los turnos de octubre y «Noviembre 2026» en la cabecera. Ese papel
    se reparte, y quien lo recibe no tiene forma de saber cuál de las dos cosas
    es la equivocada.
    """
    de_octubre = _guardar(2026, 10)

    respuesta = cliente.post('/api/exportacion/exportar/2026/11',
                             json={'libro': 'completo', 'horario_id': de_octubre})
    assert respuesta.status_code >= 400, (
        'se exportó el horario de octubre con el nombre de noviembre')
    assert 'octubre' in respuesta.text and 'noviembre' in respuesta.text


# ------------------------------------- H28 · un nombre no es una fórmula de Excel

def test_un_nombre_que_empieza_por_igual_se_imprime_tal_cual():
    """Los textos libres se entregaban a openpyxl sin mirar.

    Una persona apuntada como `=1+1` se guardaba como fórmula, y el Excel que se
    reparte enseñaba `2` donde tenía que ir su nombre. Con algo menos inocente
    —y las hojas de cálculo tienen funciones que leen archivos y abren
    enlaces— el papel deja de ser un papel.
    """
    from openpyxl import Workbook

    from gestor.servicios import excel

    hoja = Workbook().active
    for texto in ('=1+1', '+A1', '-2+3', '@SUM(A1)'):
        celda = excel._solo_texto(hoja.cell(1, 1, texto))
        assert celda.data_type == 's', f'«{texto}» se guardó como fórmula'
        assert celda.value == texto, 'y tiene que verse tal cual'


def test_un_nombre_normal_no_se_toca():
    from openpyxl import Workbook

    from gestor.servicios import excel

    hoja = Workbook().active
    celda = excel._solo_texto(hoja.cell(1, 1, 'Ximena Rocío Peralta Osorio'))
    assert celda.value == 'Ximena Rocío Peralta Osorio'
