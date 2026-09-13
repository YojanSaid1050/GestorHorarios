# -*- coding: utf-8 -*-
"""El libro que se imprime tiene que decir lo que la aplicación hace.

Un Excel es lo único de todo esto que sale de la pantalla y acaba pegado en una
pared. Si dice una cosa y la aplicación hace otra, gana el papel: la gente lee
el papel.

Esta prueba existe por un fallo que viajaba impreso. La hoja de comentarios
anunciaba «el tope configurado es de 7 jornadas seguidas» **siempre**, aunque la
regla fuera de 10 o de 12. La causa: leía el tope de un módulo del proyecto
anterior que aquí no existe, el import fallaba en cada exportación y un
`except Exception` lo dejaba caer en el número de respaldo. No daba error, no
salía en ningún registro, y el número equivocado se repartía en papel.
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from gestor.datos import horarios
from gestor.servicios import excel, generacion, reglas_operacion, siembra


@pytest.fixture
def octubre(base):
    siembra.sembrar()
    generado = generacion.generar(10, 2026)
    generacion.marcar_oficial(generado.propuestas[0]['horario_id'])
    guardado = horarios.oficial(2026, 10)
    datos = dict(guardado['datos'])
    datos['mes'], datos['anio'] = 10, 2026
    return datos


def _textos(ruta) -> list[str]:
    libro = load_workbook(ruta)
    return [str(celda) for hoja in libro.sheetnames
            for fila in libro[hoja].iter_rows(values_only=True)
            for celda in fila if celda]


def _fatiga(datos) -> str:
    textos = [t for t in _textos(excel.crear_excel(datos, 'completo'))
              if 'jornadas seguidas' in t]
    assert textos, 'la hoja ya no explica el control de fatiga'
    return textos[0]


@pytest.mark.lenta
@pytest.mark.parametrize('tope', [9, 12])
def test_el_excel_anuncia_el_tope_con_el_que_se_armó_el_mes(base, tope):
    """El número del papel es el de ese mes, no el de hoy."""
    siembra.sembrar()
    reglas_operacion.guardar('2026-08-01', tope, 'para la prueba')
    generado = generacion.generar(10, 2026)
    generacion.marcar_oficial(generado.propuestas[0]['horario_id'])
    datos = dict(horarios.oficial(2026, 10)['datos'])
    datos['mes'], datos['anio'] = 10, 2026

    assert f'{tope} jornadas seguidas' in _fatiga(datos)


@pytest.mark.lenta
def test_cambiar_la_regla_no_reescribe_el_papel_de_un_mes_ya_hecho(octubre):
    """Es el mismo principio que sostiene toda la aplicación.

    Un mes se mide —y se imprime— con las reglas con las que se hizo. Si el
    Excel de octubre cambiara de número porque en diciembre se subió el tope,
    el papel de la pared dejaría de corresponderse con el mes que describe.
    """
    antes = _fatiga(octubre)
    reglas_operacion.guardar('2026-12-07', 14, 'cambio posterior')
    assert _fatiga(octubre) == antes


@pytest.mark.lenta
def test_un_mes_sin_esa_anotación_lee_la_regla_en_vez_de_inventarse_un_7(base):
    """El caso exacto del fallo: los meses base no traen esa anotación.

    Agosto y septiembre llegan transcritos y no guardan con qué tope se
    hicieron. Ahí es donde saltaba el respaldo, y donde el papel decía 7.
    """
    siembra.sembrar()
    reglas_operacion.guardar('2026-08-01', 11, 'la que rige')
    guardado = horarios.oficial(2026, 9)
    datos = dict(guardado['datos'])
    datos['mes'], datos['anio'] = 9, 2026
    assert 'maximo_dias_consecutivos' not in (datos.get('reglas') or {})

    texto = _fatiga(datos)
    assert '11 jornadas seguidas' in texto, texto
    assert '7 jornadas seguidas' not in texto


@pytest.mark.lenta
def test_el_excel_no_se_inventa_un_tope_de_respaldo(octubre):
    """Y si no hubiera regla configurada, el respaldo es el del propio módulo.

    No un 7 escrito a mano en otro archivo: dos números para la misma cosa es
    como empiezan a decir cosas distintas.
    """
    assert excel.MAX_DIAS_POR_DEFECTO == reglas_operacion.POR_DEFECTO


@pytest.mark.lenta
@pytest.mark.parametrize('libro', ['trabajo', 'semana', 'completo'])
def test_los_tres_libros_salen(octubre, libro):
    if libro == 'semana':
        octubre['_export_week_start'] = octubre['horario'][0]['dias'][0]['lunes_semana']
    ruta = excel.crear_excel(octubre, libro)
    assert ruta.is_file() and ruta.stat().st_size > 5000, ruta


@pytest.mark.lenta
def test_el_libro_de_trabajo_lleva_sus_dos_hojas(octubre):
    """Se cambia un turno en «Semanas completas» y «Solo el mes» lo recoge sola."""
    libro = load_workbook(excel.crear_excel(octubre, 'trabajo'))
    assert libro.sheetnames == ['Semanas completas', 'Solo el mes'], libro.sheetnames


@pytest.mark.lenta
def test_un_horario_vacio_se_niega_en_vez_de_sacar_un_libro_en_blanco(base):
    with pytest.raises(ValueError, match='vacío'):
        excel.crear_excel({'horario': [], 'mes': 10, 'anio': 2026}, 'completo')
